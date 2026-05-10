import sys
import subprocess
import os

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Pt, Inches

def create_guide():
    document = Document()
    
    # Add Title
    title = document.add_heading('StockSense (Stock_LSTM) Installation and Run Guide', 0)
    
    document.add_heading('1. Prerequisites', level=1)
    document.add_paragraph('Before you begin, ensure you have the following installed on your system:')
    p = document.add_paragraph('', style='List Bullet')
    p.add_run('Python: ').bold = True
    p.add_run('Version 3.8 or above is recommended.')
    p = document.add_paragraph('', style='List Bullet')
    p.add_run('Git: ').bold = True
    p.add_run('Optional, if you need to clone the repository.')
    
    document.add_heading('2. Setup and Installation', level=1)
    document.add_heading('Step 2.1: Open Terminal or Command Prompt', level=2)
    document.add_paragraph('Navigate to the project directory where your files are located:')
    document.add_paragraph('cd path\\to\\Stock_LSTM', style='Intense Quote')
    
    document.add_heading('Step 2.2: Create a Virtual Environment', level=2)
    document.add_paragraph('It is highly recommended to use a virtual environment to avoid dependency conflicts. Run the following command:')
    document.add_paragraph('python -m venv venv', style='Intense Quote')
    
    document.add_heading('Step 2.3: Activate the Virtual Environment', level=2)
    p = document.add_paragraph('Depending on your operating system, run the activation command:')
    document.add_paragraph('Windows: .\\venv\\Scripts\\activate', style='Intense Quote')
    document.add_paragraph('macOS/Linux: source venv/bin/activate', style='Intense Quote')
    
    document.add_heading('Step 2.4: Install Dependencies', level=2)
    document.add_paragraph('Install all required libraries using the requirements.txt file:')
    document.add_paragraph('pip install -r requirements.txt', style='Intense Quote')
    document.add_paragraph('This will install packages like Flask, TensorFlow, Keras, scikit-learn, etc.')
    
    document.add_heading('3. Running the Application', level=1)
    
    document.add_heading('Step 3.1: Start the Server', level=2)
    document.add_paragraph('Once the dependencies are installed, start the Flask web application by running:')
    document.add_paragraph('python app.py', style='Intense Quote')
    
    document.add_heading('Step 3.2: Access the Application', level=2)
    document.add_paragraph('Open your web browser and navigate to the following URL:')
    document.add_paragraph('http://127.0.0.1:5000', style='Intense Quote')
    
    document.add_heading('4. Optional Configuration', level=1)
    document.add_paragraph('The application can work without API keys by degrading gracefully to fallbacks, but you can configure the following environment variables if needed:')
    p = document.add_paragraph('', style='List Bullet')
    p.add_run('GNEWS_API_KEY: ').bold = True
    p.add_run('Use an API key from GNews to fetch live headlines instead of falling back to Yahoo Finance.')
    p = document.add_paragraph('', style='List Bullet')
    p.add_run('FLASK_SECRET_KEY: ').bold = True
    p.add_run('Used for securing the session in user authentication.')
    p = document.add_paragraph('', style='List Bullet')
    p.add_run('DATABASE_URL: ').bold = True
    p.add_run('Custom database URI (by default it uses SQLite).')
    
    document.add_heading('5. Troubleshooting', level=1)
    p = document.add_paragraph('', style='List Bullet')
    p.add_run('Database errors: ').bold = True
    p.add_run('The project defaults to SQLite and creates the dataset database "instance/users.db" automatically. ')
    p = document.add_paragraph('', style='List Bullet')
    p.add_run('Health Check: ').bold = True
    p.add_run('You can verify the app status by visiting http://127.0.0.1:5000/health in your browser.')
    
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'StockSense_Installation_Guide.docx')
    document.save(path)
    print(f"Created documentation successfully: {path}")

if __name__ == "__main__":
    create_guide()
