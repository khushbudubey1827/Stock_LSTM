from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pathlib import Path

ROOT = Path(r"c:\Users\ASUS\Desktop\Stock_LSTM")
OUT = ROOT / "StockSense_LSTM_Project_Presentation.pptx"

prs = Presentation()


def set_bg(slide, color=(245, 248, 252)):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*color)


def add_title_slide(title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    set_bg(slide, (236, 244, 255))
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullet_slide(title, bullets, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_bg(slide)
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()

    if subtitle:
        p = body.paragraphs[0]
        p.text = subtitle
        p.font.bold = True
        p.font.size = Pt(20)

    for b in bullets:
        p = body.add_paragraph() if body.paragraphs[0].text else body.paragraphs[0]
        p.text = b
        p.level = 0
        p.font.size = Pt(20)


def add_image_slide(title, image_path, caption):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    set_bg(slide)
    slide.shapes.title.text = title

    left = Inches(0.6)
    top = Inches(1.2)
    width = Inches(12.1)
    height = Inches(5.5)

    slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)

    cap = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.5), Inches(0.5)).text_frame
    cap.text = caption
    cap.paragraphs[0].font.size = Pt(16)
    cap.paragraphs[0].font.color.rgb = RGBColor(60, 60, 60)


def add_two_col_slide(title, left_points, right_points, left_header, right_header):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    set_bg(slide)
    slide.shapes.title.text = title

    left_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(5.9), Inches(5.6)).text_frame
    right_box = slide.shapes.add_textbox(Inches(6.8), Inches(1.4), Inches(5.9), Inches(5.6)).text_frame

    left_box.text = left_header
    left_box.paragraphs[0].font.bold = True
    left_box.paragraphs[0].font.size = Pt(24)
    for item in left_points:
        p = left_box.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(18)

    right_box.text = right_header
    right_box.paragraphs[0].font.bold = True
    right_box.paragraphs[0].font.size = Pt(24)
    for item in right_points:
        p = right_box.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(18)


add_title_slide(
    "StockSense: LSTM-Based Stock Prediction Platform",
    "Project Presentation | Flask + Deep Learning + Financial Analytics"
)

add_bullet_slide(
    "Problem Statement",
    [
        "Retail traders often lack simple, integrated tools for stock analysis.",
        "Most solutions separate technical charts, prediction, and sentiment.",
        "Goal: build one web app that gives trend insight, next-day estimate, and market mood."
    ]
)

add_bullet_slide(
    "Project Overview",
    [
        "Backend: Flask application with route-based modules for dashboard, live price, prediction, and sentiment.",
        "Model: pre-trained deep LSTM network loaded from stock_dl_model.h5.",
        "Data sources: Yahoo Finance (historical + live), GNews/YFinance headlines for sentiment.",
        "Output: charts, next-day signal (BUY/SELL/HOLD), downloadable datasets, and live metrics."
    ]
)

add_two_col_slide(
    "System Architecture",
    left_points=[
        "User authentication (register/login/logout)",
        "Dashboard route (/) for model visualization",
        "Live price API and page",
        "Next-day prediction page",
        "Sentiment analysis page",
        "CSV dataset download endpoint"
    ],
    right_points=[
        "MySQL + SQLAlchemy for user records",
        "Keras model inference for close-price trend",
        "MinMaxScaler for normalization/inverse transform",
        "Matplotlib chart generation",
        "TextBlob polarity scoring",
        "Session-based protected routes"
    ],
    left_header="Application Layer",
    right_header="ML + Data Layer"
)

add_bullet_slide(
    "Model Training Pipeline (Notebook)",
    [
        "Example ticker used: POWERGRID.NS with 4,555 rows of historical OHLCV data.",
        "Data split: 70% training, 30% testing on Close price.",
        "Sequence setup: 100-day lookback window for each sample.",
        "Scaling: MinMaxScaler (0 to 1) for stable LSTM learning.",
        "Training: 50 epochs, Adam optimizer, MSE loss."
    ]
)

add_bullet_slide(
    "LSTM Network Configuration",
    [
        "Input shape: (100, 1)",
        "Stacked layers: LSTM(50) -> LSTM(60) -> LSTM(80) -> LSTM(120)",
        "Dropout: 0.2, 0.3, 0.4, 0.5 to reduce overfitting",
        "Output layer: Dense(1) for next value regression",
        "Approximate trainable parameters: 178,761"
    ]
)

add_image_slide(
    "Technical Indicator View",
    ROOT / "static" / "ema_20_50.png",
    "Closing price with EMA-20 and EMA-50 overlays generated by the Flask dashboard."
)

add_image_slide(
    "Long-Term Trend View",
    ROOT / "static" / "ema_100_200.png",
    "Closing price with EMA-100 and EMA-200 to capture long-range movement."
)

add_image_slide(
    "Prediction vs Actual",
    ROOT / "static" / "stock_prediction.png",
    "Model output compared with actual test-series movement."
)

add_two_col_slide(
    "Key Features Delivered",
    left_points=[
        "Secure user login and registration",
        "Interactive dashboard for selected ticker",
        "Auto-generated analytics charts",
        "Downloadable historical CSV datasets"
    ],
    right_points=[
        "Next-day price estimate with direction",
        "Trading signal logic: BUY / HOLD / SELL",
        "Real-time price snapshot and intraday series",
        "Headline sentiment summary with polarity labels"
    ],
    left_header="Product Features",
    right_header="Decision Support"
)

add_bullet_slide(
    "Current Limitations & Future Scope",
    [
        "Predictions are based primarily on historical close prices (single-feature focus).",
        "Sentiment uses lexical polarity; finance-specific NLP can improve reliability.",
        "Backtesting and risk metrics can be added for stronger validation.",
        "Future work: multi-feature modeling, transformer-based forecasting, portfolio-level recommendations."
    ]
)

add_bullet_slide(
    "Conclusion",
    [
        "This project demonstrates an end-to-end stock analytics platform with deployed ML inference.",
        "It combines data engineering, deep learning, and web product development in one workflow.",
        "The architecture is ready for iterative upgrades toward production-grade quant intelligence."
    ]
)

prs.save(str(OUT))
print(f"Created: {OUT}")
